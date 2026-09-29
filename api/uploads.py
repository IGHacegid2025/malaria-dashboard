# Author: Khadim Gueye

import csv
import io
import re
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from db import query

STATE_ALIASES = {
    "fct": "FC",
    "abuja": "FC",
    "fct abuja": "FC",
    "abuja fct": "FC",
    "federal capital territory": "FC",
    "federal capital territory abuja": "FC",
    "ibadan": "OY",
    "nassarawa": "NA",
    "cross rivers": "CR",
}

SPECIES_ALIASES = {
    "pf": "pf",
    "p falciparum": "pf",
    "plasmodium falciparum": "pf",
    "falciparum": "pf",
    "pm": "pm",
    "p malariae": "pm",
    "plasmodium malariae": "pm",
    "malariae": "pm",
    "po": "po",
    "p ovale": "po",
    "plasmodium ovale": "po",
    "ovale": "po",
    "pv": "pv",
    "p vivax": "pv",
    "plasmodium vivax": "pv",
    "vivax": "pv",
    "pk": "pk",
    "p knowlesi": "pk",
    "plasmodium knowlesi": "pk",
    "knowlesi": "pk",
}

GENE_ALIASES = {"k13": "kelch13", "kelch 13": "kelch13", "pfk13": "kelch13"}

DELETION_ALIASES = {
    "hrp2": "hrp2",
    "pfhrp2": "hrp2",
    "hrp3": "hrp3",
    "pfhrp3": "hrp3",
    "dual": "dual",
    "both": "dual",
    "hrp2 hrp3": "dual",
    "hrp2 and hrp3": "dual",
    "none": "none",
    "no": "none",
    "no deletion": "none",
    "negative": "none",
}


def norm_key(value) -> str:
    text = re.sub(r"[\-_./(),+&]+", " ", str(value or "").lower())
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"^state of ", "", text)
    text = re.sub(r" state$", "", text)
    return text


def norm_gene(value, known) -> str:
    text = re.sub(r"\s+", "", str(value or "").strip().lower())
    text = GENE_ALIASES.get(text, text)
    if text not in known and text.startswith("pf") and text[2:] in known:
        text = text[2:]
    return GENE_ALIASES.get(text, text)

SOURCES = ["Sequencing", "Publication"]
DELETION_TYPES = ["hrp2", "hrp3", "dual", "none"]

PUBLICATION_COLUMNS = [
    ("first_author", "First author", False, "Surname of the first author, e.g. Beshir. Required when Source is Publication."),
    ("publication_year", "Publication year", False, "Year the paper was published, e.g. 2023. Required when Source is Publication."),
    ("doi", "DOI", False, "Digital Object Identifier, e.g. 10.1186/s12936-020-03506-z. Used to avoid duplicate publications."),
    ("title", "Title", False, "Title of the publication (optional)."),
]

KINDS = {
    "sequencing_mutations": {
        "label": "Sequencing: resistance mutations",
        "description": "Prevalence of drug resistance mutations measured by IGH sequencing, one row per state, year and mutation.",
        "columns": [
            ("year", "Year", True, "Year of sample collection, e.g. 2021."),
            ("state", "State", True, "Nigerian state name or 2-letter code (see the Lists sheet)."),
            ("gene", "Gene", True, "Gene name, e.g. kelch13, crt, mdr1, dhfr, dhps, psfr, coronin."),
            ("mutation", "Mutation", True, "Mutation or haplotype code, e.g. C580Y or N86Y-Y184F-D1246N."),
            ("prevalence", "Prevalence", True, "Share of samples carrying the mutation: 0 to 1 (0.25) or percent (25%)."),
            ("sample_count", "Sample count", True, "Number of samples genotyped for this mutation."),
        ],
        "example": [2021, "Kano", "kelch13", "C580Y", 0.012, 250],
    },
    "publication_mutations": {
        "label": "Publications: resistance mutations",
        "description": "Mutation prevalences reported in published studies, one row per study, state, year and mutation.",
        "columns": [
            ("first_author", "First author", True, "Surname of the first author, e.g. Martin Ramirez."),
            ("publication_year", "Publication year", True, "Year the paper was published, e.g. 2025."),
            ("doi", "DOI", False, "Digital Object Identifier. Strongly recommended to avoid duplicates."),
            ("title", "Title", False, "Title of the publication."),
            ("year", "Sample collection year", True, "Year the samples were collected (corrected year), e.g. 2021."),
            ("state", "State", True, "Nigerian state name or 2-letter code."),
            ("city", "City", False, "City or site of sample collection."),
            ("gene", "Gene", True, "Gene name, e.g. kelch13, crt, mdr1."),
            ("mutation", "Mutation", True, "Mutation or haplotype code, e.g. C580Y."),
            ("prevalence", "Prevalence", True, "0 to 1 (0.0123) or percent (1.23%)."),
            ("sample_count", "Sample count", False, "Number of samples tested in the study for this state."),
        ],
        "example": ["Martin Ramirez", 2025, "10.1186/s41182-025-00732-6", "Mutational profile of pfdhfr, pfdhps, pfmdr1, pfcrt and pfk13", 2021, "Osun", "Ore", "kelch13", "C580Y", "1.23%", 262],
    },
    "species": {
        "label": "Plasmodium species",
        "description": "Number of samples by Plasmodium species or mixed infection.",
        "columns": [
            ("year", "Year", True, "Year of sample collection."),
            ("state", "State", True, "Nigerian state name or 2-letter code."),
            ("source", "Source", True, "Sequencing or Publication."),
            *PUBLICATION_COLUMNS,
            ("species", "Species", True, "One species (pf) or a mixed infection joined with + (pf+pm). Names like P. falciparum also work."),
            ("sample_count", "Sample count", True, "Number of samples with this species result."),
        ],
        "example": [2021, "Bauchi", "Sequencing", "", "", "", "", "pf+pm", 12],
    },
    "diagnostic": {
        "label": "Diagnostics: hrp2 / hrp3 deletions",
        "description": "Samples tested for hrp2/hrp3 gene deletions. Use one row per deletion type, including none.",
        "columns": [
            ("year", "Year", True, "Year of sample collection."),
            ("state", "State", True, "Nigerian state name or 2-letter code."),
            ("source", "Source", True, "Sequencing or Publication."),
            *PUBLICATION_COLUMNS,
            ("location", "Location", False, "Site or health facility (optional)."),
            ("deletion_type", "Deletion type", True, "hrp2, hrp3, dual (both deleted) or none (no deletion)."),
            ("sample_count", "Sample count", True, "Number of samples with this result."),
        ],
        "example": [2021, "Borno", "Sequencing", "", "", "", "", "", "hrp2", 3],
    },
    "moi": {
        "label": "Parasite clones (MOI)",
        "description": "Distribution of the multiplicity of infection from sequencing: number of samples per number of clones.",
        "columns": [
            ("year", "Year", True, "Year of sample collection."),
            ("state", "State", True, "Nigerian state name or 2-letter code."),
            ("moi", "Clones (MOI)", True, "Number of distinct clones in the infection: 0 (no call), 1, 2, 3..."),
            ("sample_count", "Sample count", True, "Number of samples with this MOI."),
        ],
        "example": [2021, "Lagos", 2, 7],
    },
    "mis": {
        "label": "Malaria in children (MIS)",
        "description": "Share of children testing positive for malaria in a national survey, one row per state and year.",
        "columns": [
            ("year", "Year", True, "Survey year, e.g. 2021."),
            ("state", "State", True, "Nigerian state name or 2-letter code."),
            ("prevalence", "Prevalence", True, "0 to 1 (0.188) or percent (18.8%)."),
        ],
        "example": [2021, "Kebbi", "49%"],
    },
    "who_thresholds": {
        "label": "WHO thresholds and alert rules",
        "description": "Alert levels, messages and WHO status per marker. Same columns as information_for_report.xlsx; that file can be uploaded as it is.",
        "columns": [
            ("category", "Category", False, "drug resistance or diagnostic resistance."),
            ("antimalarial", "Antimalarial", False, "Drug affected, e.g. artemisinin. Use NA for diagnostics."),
            ("gene", "Marker", True, "Gene, e.g. kelch13, crt, mdr1, dhps/dhfr, hrp2."),
            ("mutation", "Haplotype", True, "Mutation or haplotype (C580Y, N86Y, Y184F, D1246N). Write 'any additional mutation' for the general rule of a gene, 'deletion' for hrp2/hrp3."),
            ("who_status", "WHO status", False, "validated, candidate or none. Left empty, it is read from the message."),
            ("min", "Min", False, "Lower bound of the level, 0 to 1 (e.g. 0.31)."),
            ("max", "Max", True, "Upper bound of the level, 0 to 1 (e.g. 0.6). One row per level."),
            ("classification", "Classification", True, "low, intermediate or high."),
            ("message", "Message", False, "Detailed message shown on hover."),
            ("summary", "Action", False, "Short action text shown on the cards."),
            ("comment", "Comment", False, "Reference or note, e.g. WHO guideline link."),
        ],
        "example": ["drug resistance", "artemisinin", "kelch13", "C580Y", "validated", 0, 1, "high",
                    "This molecular marker is a validated mutation leading to artemisinin partial resistance.",
                    "Reduced sensitivity to artemisinin is expected.", "https://www.who.int/"],
    },
}

HEADER_ALIASES = {"marker": "marker", "gene": "marker", "haplotype": "haplotype", "mutation": "haplotype",
                  "comment": "comment", "reference": "comment", "action": "action", "summary": "action"}
HEADER_FILL = PatternFill("solid", fgColor="0D366B")
REQUIRED_FILL = PatternFill("solid", fgColor="184F95")


def reference_lists():
    states = query("SELECT code, name FROM states ORDER BY name")
    genes = [g["name"] for g in query("SELECT name FROM genes ORDER BY name")]
    rows = query(
        """
        SELECT g.name AS gene, m.mutation_code
        FROM mutations m JOIN genes g ON g.id = m.gene_id
        ORDER BY g.name, m.mutation_code
        """
    )
    mutations = {g: [] for g in genes}
    for r in rows:
        mutations[r["gene"]].append(r["mutation_code"])
    return states, genes, mutations


SPECIES_CHOICES = ["pf", "pm", "po", "pv", "pk", "pf+pm", "pf+po", "pm+po", "pf+pm+po"]


def build_template(kind: str) -> bytes:
    spec = KINDS[kind]
    states, genes, mutations = reference_lists()
    wb = Workbook()

    info = wb.active
    info.title = "Instructions"
    info["A1"] = spec["label"]
    info["A1"].font = Font(size=16, bold=True, color="0D366B")
    info["A2"] = spec["description"]
    info["A4"] = "How to use this file"
    info["A4"].font = Font(bold=True)
    steps = [
        "1. Fill the Data sheet, one row per record. Do not rename or move the column headers.",
        "2. Columns marked with * are required. Light blue columns have a dropdown: click a cell, then the arrow on its right.",
        "3. For mutations, choose the Gene first: the Mutation dropdown then lists the known mutations of that gene. A new code can still be typed.",
        "4. See the Example sheet for a filled row. The Lists sheet shows accepted states, genes, species and mutations.",
        "5. Save the file and upload it in the admin area (Submit data). You will see a check of every row before it is saved.",
    ]
    for i, step in enumerate(steps, start=5):
        info.cell(row=i, column=1, value=step)
    info.cell(row=10, column=1, value="Column").font = Font(bold=True)
    info.cell(row=10, column=2, value="Required").font = Font(bold=True)
    info.cell(row=10, column=3, value="Description").font = Font(bold=True)
    for i, (_, header, required, help_text) in enumerate(spec["columns"], start=11):
        info.cell(row=i, column=1, value=header)
        info.cell(row=i, column=2, value="Yes" if required else "No")
        info.cell(row=i, column=3, value=help_text)
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 10
    info.column_dimensions["C"].width = 110

    data = wb.create_sheet("Data")
    example = wb.create_sheet("Example")
    lists = wb.create_sheet("Lists")

    for sheet in (data, example):
        for col, (_, header, required, help_text) in enumerate(spec["columns"], start=1):
            cell = sheet.cell(row=1, column=col, value=f"{header} *" if required else header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = REQUIRED_FILL if required else HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(col)].width = max(14, min(40, len(header) + 8))
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30
    for col, value in enumerate(spec["example"], start=1):
        example.cell(row=2, column=col, value=value)

    lists["A1"], lists["B1"], lists["C1"], lists["D1"], lists["E1"] = "State", "Code", "Gene", "Source", "Deletion type"
    for c in "ABCDE":
        lists[f"{c}1"].font = Font(bold=True)
    for i, s in enumerate(states, start=2):
        lists.cell(row=i, column=1, value=s["name"])
        lists.cell(row=i, column=2, value=s["code"])
    for i, g in enumerate(genes, start=2):
        lists.cell(row=i, column=3, value=g)
    for i, s in enumerate(SOURCES, start=2):
        lists.cell(row=i, column=4, value=s)
    for i, d in enumerate(DELETION_TYPES, start=2):
        lists.cell(row=i, column=5, value=d)
    lists["F1"] = "Species"
    lists["F1"].font = Font(bold=True)
    for i, s in enumerate(SPECIES_CHOICES, start=2):
        lists.cell(row=i, column=6, value=s)
    for c in "ABCDEF":
        lists.column_dimensions[c].width = 26

    named = {
        "StateList": f"Lists!$A$2:$A${len(states) + 1}",
        "GeneList": f"Lists!$C$2:$C${len(genes) + 1}",
        "SourceList": f"Lists!$D$2:$D${len(SOURCES) + 1}",
        "DeletionList": f"Lists!$E$2:$E${len(DELETION_TYPES) + 1}",
        "SpeciesList": f"Lists!$F$2:$F${len(SPECIES_CHOICES) + 1}",
    }
    lists.cell(row=1, column=8, value="Mutations by gene (used by the Mutation dropdown)").font = Font(bold=True)
    for offset, gene in enumerate(genes):
        col_index = 8 + offset
        col_letter = get_column_letter(col_index)
        lists.cell(row=2, column=col_index, value=gene).font = Font(bold=True, color="0D366B")
        codes = mutations.get(gene) or [""]
        for i, code in enumerate(codes, start=3):
            lists.cell(row=i, column=col_index, value=code)
        lists.column_dimensions[col_letter].width = max(14, min(42, max(len(c) for c in codes) + 4))
        named[f"Mut_{gene}"] = f"Lists!${col_letter}$3:${col_letter}${len(codes) + 2}"
    for name, ref in named.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)

    keys = [c[0] for c in spec["columns"]]
    last_year = date.today().year + 1
    rules = {
        "state": dict(type="list", formula1="StateList", promptTitle="State",
                      prompt="Click the arrow and choose a state. Upper or lower case both work.",
                      errorTitle="State not in the list", error="Choose a state from the list, or keep your value if it is a code or known alias."),
        "gene": dict(type="list", formula1="GeneList", promptTitle="Gene",
                     prompt="Choose a gene from the list, or type a new one.",
                     errorTitle="New gene", error="This gene is not in the list yet. Keep it only if it is a new gene."),
        "source": dict(type="list", formula1="SourceList", promptTitle="Source",
                       prompt="Choose Sequencing or Publication.",
                       errorTitle="Source", error="Choose Sequencing or Publication."),
        "deletion_type": dict(type="list", formula1="DeletionList", promptTitle="Deletion type",
                              prompt="Choose hrp2, hrp3, dual (both deleted) or none.",
                              errorTitle="Deletion type", error="Choose hrp2, hrp3, dual or none."),
        "who_status": dict(type="list", formula1='"validated,candidate,none"', promptTitle="WHO status",
                           prompt="validated: WHO-validated marker. candidate: WHO candidate marker. none: team judgement.",
                           errorTitle="WHO status", error="Use validated, candidate or none."),
        "classification": dict(type="list", formula1='"low,intermediate,high"', promptTitle="Classification",
                               prompt="Alert level for this range.", errorTitle="Classification", error="Use low, intermediate or high."),
        "species": dict(type="list", formula1="SpeciesList", promptTitle="Species",
                        prompt="Choose one species (pf) or a mixed infection (pf+pm).",
                        errorTitle="Species", error="Use pf, pm, po, pv, pk, or a mix joined with + (pf+pm)."),
        "year": dict(type="whole", operator="between", formula1="1980", formula2=str(last_year), promptTitle="Year",
                     prompt="Four-digit year, e.g. 2021.", errorTitle="Year", error=f"Enter a year between 1980 and {last_year}."),
        "publication_year": dict(type="whole", operator="between", formula1="1980", formula2=str(last_year),
                                 promptTitle="Publication year", prompt="Four-digit year, e.g. 2023.",
                                 errorTitle="Year", error=f"Enter a year between 1980 and {last_year}."),
        "sample_count": dict(type="whole", operator="greaterThanOrEqual", formula1="0", promptTitle="Sample count",
                             prompt="Whole number of samples.", errorTitle="Sample count", error="Enter a whole number, 0 or more."),
        "moi": dict(type="whole", operator="between", formula1="0", formula2="50", promptTitle="Clones (MOI)",
                    prompt="Number of clones: 0 (no call), 1, 2, 3...", errorTitle="MOI", error="Enter a whole number from 0 to 50."),
        "prevalence": dict(type="custom", formula1="TRUE", promptTitle="Prevalence",
                           prompt="0 to 1 (0.25) or a percentage (25%)."),
    }
    if "mutation" in keys and "gene" in keys:
        gene_col = get_column_letter(keys.index("gene") + 1)
        rules["mutation"] = dict(
            type="list",
            formula1=f'INDIRECT("Mut_"&${gene_col}2)',
            promptTitle="Mutation",
            prompt="Choose the gene first: this list then shows its known mutations. You can also type a new code.",
            errorTitle="New mutation",
            error="This mutation is not in the list for this gene. Keep it only if it is a new mutation.",
        )
    list_fill = PatternFill("solid", fgColor="EEF4FC")
    for key, options in rules.items():
        if key not in keys:
            continue
        col = get_column_letter(keys.index(key) + 1)
        dv = DataValidation(allow_blank=True, showInputMessage=True, showErrorMessage=True, errorStyle="warning", **options)
        dv.add(f"{col}2:{col}5000")
        data.add_data_validation(dv)
        if options["type"] == "list":
            for r in range(2, 201):
                data[f"{col}{r}"].fill = list_fill

    wb.active = 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _norm_header(text) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").replace("*", "").lower()).strip()


def read_rows(filename: str, content: bytes):
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".tsv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        delimiter = "\t" if name.endswith(".tsv") or text.count("\t") > text.count(",") else ","
        rows = [r for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    else:
        raise ValueError("Unsupported file type. Use .xlsx, .csv or .tsv.")
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        raise ValueError("The file is empty.")
    return rows[0], rows[1:]


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _int(value):
    text = _text(value)
    if text == "":
        return None
    try:
        number = float(text)
    except ValueError:
        raise ValueError(f"'{text}' is not a number")
    if not number.is_integer():
        raise ValueError(f"'{text}' must be a whole number")
    return int(number)


def _prevalence(value):
    text = _text(value)
    if text == "":
        return None, None
    percent = text.endswith("%")
    try:
        number = float(text.rstrip("%").replace(",", "."))
    except ValueError:
        raise ValueError(f"'{text}' is not a number")
    warning = None
    if percent:
        number /= 100
    elif number > 1:
        if number <= 100:
            warning = f"{text} read as {number}%"
            number /= 100
        else:
            raise ValueError(f"{text} is above 100%")
    if number < 0 or number > 1:
        raise ValueError(f"'{text}' must be between 0 and 1 (or 0% to 100%)")
    return round(number, 6), warning


def normalise_doi(value):
    text = _text(value)
    return re.sub(r"^(https?://)?(dx\.)?doi\.org/", "", text, flags=re.IGNORECASE) or None


def validate(kind: str, filename: str, content: bytes):
    if kind == "who_thresholds":
        return validate_rules(filename, content)
    spec = KINDS[kind]
    header, body = read_rows(filename, content)
    states = query("SELECT id, code, name FROM states")
    by_key = {}
    for s in states:
        by_key[norm_key(s["code"])] = s
        by_key[norm_key(s["name"])] = s
        by_key[norm_key(s["name"]).replace(" ", "")] = s
    for alias, code in STATE_ALIASES.items():
        by_key[alias] = next(s for s in states if s["code"] == code)
    known_genes = {g["name"].lower() for g in query("SELECT name FROM genes")}

    wanted = {key: _norm_header(label) for key, label, _, _ in spec["columns"]}
    header_norm = [_norm_header(h) for h in header]
    positions = {}
    missing = []
    for key, label in wanted.items():
        if label in header_norm:
            positions[key] = header_norm.index(label)
        elif any(c[0] == key and c[2] for c in spec["columns"]):
            missing.append(next(c[1] for c in spec["columns"] if c[0] == key))
    if missing:
        return {
            "total_rows": len(body),
            "valid_rows": 0,
            "errors": [{"row": 1, "column": "Header", "message": f"Missing column(s): {', '.join(missing)}. Use the template headers."}],
            "warnings": [],
            "rows": [],
        }

    errors, warnings, clean = [], [], []
    seen = set()
    this_year = date.today().year
    for index, raw in enumerate(body, start=2):
        values = {key: (raw[pos] if pos < len(raw) else None) for key, pos in positions.items()}
        row = {}
        row_errors = []

        def err(column, message):
            row_errors.append({"row": index, "column": column, "message": message})

        for key, label, required, _ in spec["columns"]:
            if required and _text(values.get(key)) == "":
                err(label, "Required value is empty")

        try:
            year = _int(values.get("year"))
            if year is not None and not 1980 <= year <= this_year + 1:
                err("Year", f"{year} is not a plausible year")
            row["year"] = year
        except ValueError as e:
            err("Year", str(e))

        state_text = norm_key(_text(values.get("state")))
        if state_text:
            state = by_key.get(state_text) or by_key.get(state_text.replace(" ", ""))
            if state:
                row["state_id"], row["state_code"], row["state_name"] = state["id"], state["code"], state["name"]
            else:
                err("State", f"Unknown state '{_text(values.get('state'))}'")

        if "gene" in positions:
            gene = norm_gene(_text(values.get("gene")), known_genes)
            row["gene"] = gene
            if gene and gene not in known_genes:
                warnings.append({"row": index, "column": "Gene", "message": f"New gene '{gene}' will be created"})
        if "mutation" in positions:
            row["mutation"] = re.sub(r"\s+", "", _text(values.get("mutation"))).upper()

        if "prevalence" in positions:
            try:
                row["prevalence"], note = _prevalence(values.get("prevalence"))
                if note:
                    warnings.append({"row": index, "column": "Prevalence", "message": note})
            except ValueError as e:
                err("Prevalence", str(e))

        for key, label in (("sample_count", "Sample count"), ("moi", "Clones (MOI)"), ("publication_year", "Publication year")):
            if key in positions:
                try:
                    number = _int(values.get(key))
                    if number is not None and number < 0:
                        err(label, "Must be zero or more")
                    row[key] = number
                except ValueError as e:
                    err(label, str(e))

        source = "publication" if kind == "publication_mutations" else None
        if "source" in positions:
            source_text = norm_key(_text(values.get("source")))
            if source_text.startswith("seq") or "sequenc" in source_text:
                source = "sequencing"
            elif source_text.startswith("pub") or source_text in ("paper", "literature", "study"):
                source = "publication"
            elif source_text:
                err("Source", "Use Sequencing or Publication")
        if kind in ("sequencing_mutations", "moi"):
            source = "sequencing"
        row["source"] = source

        if source == "publication":
            row["first_author"] = _text(values.get("first_author"))
            row["doi"] = normalise_doi(values.get("doi"))
            row["title"] = _text(values.get("title")) or None
            if not row["first_author"]:
                err("First author", "Required for publication data")
            if row.get("publication_year") is None:
                err("Publication year", "Required for publication data")

        if "city" in positions:
            row["city"] = _text(values.get("city")) or None
        if "location" in positions:
            row["location"] = _text(values.get("location")) or None

        if "species" in positions:
            parts = [norm_key(p) for p in re.split(r"[+,;/&]| and ", _text(values.get("species")).lower()) if p.strip()]
            codes = []
            for p in parts:
                code = SPECIES_ALIASES.get(p)
                if code:
                    codes.append(code)
                else:
                    err("Species", f"Unknown species '{p}'")
            row["species"] = ",".join(sorted(set(codes), key=codes.index))

        if "deletion_type" in positions:
            raw_deletion = norm_key(_text(values.get("deletion_type")))
            deletion = DELETION_ALIASES.get(raw_deletion, DELETION_ALIASES.get(raw_deletion.replace(" ", ""), ""))
            if raw_deletion and not deletion:
                err("Deletion type", "Use hrp2, hrp3, dual or none")
            row["deletion_type"] = deletion

        if row_errors:
            errors.extend(row_errors)
            continue

        signature = tuple(sorted((k, str(v)) for k, v in row.items()))
        if signature in seen:
            warnings.append({"row": index, "column": "", "message": "Same values as an earlier row"})
        seen.add(signature)
        row["_row"] = index
        clean.append(row)

    return {
        "total_rows": len(body),
        "valid_rows": len(clean),
        "errors": errors,
        "warnings": warnings,
        "rows": clean,
    }


def _gene_id(cur, name):
    cur.execute("SELECT id FROM genes WHERE name = %s", (name,))
    found = cur.fetchone()
    if found:
        return found["id"]
    cur.execute("INSERT INTO genes (name) VALUES (%s)", (name,))
    return cur.lastrowid


def _mutation_id(cur, gene, code):
    gene_id = _gene_id(cur, gene)
    cur.execute("SELECT id FROM mutations WHERE gene_id = %s AND mutation_code = %s", (gene_id, code))
    found = cur.fetchone()
    if found:
        return found["id"]
    cur.execute("INSERT INTO mutations (gene_id, mutation_code) VALUES (%s, %s)", (gene_id, code))
    return cur.lastrowid


def _batch_id(cur, state_id, year):
    cur.execute("SELECT id FROM sequencing_batches WHERE state_id = %s AND year = %s", (state_id, year))
    found = cur.fetchone()
    if found:
        return found["id"]
    cur.execute("INSERT INTO sequencing_batches (state_id, year) VALUES (%s, %s)", (state_id, year))
    return cur.lastrowid


def _publication_id(cur, row, upload_id):
    if row.get("doi"):
        cur.execute("SELECT id FROM publications WHERE doi = %s", (row["doi"],))
    else:
        cur.execute(
            "SELECT id FROM publications WHERE author = %s AND year_of_publication = %s AND doi IS NULL",
            (row["first_author"], row["publication_year"]),
        )
    found = cur.fetchone()
    if found:
        return found["id"]
    cur.execute(
        "INSERT INTO publications (author, year_of_publication, doi, title, upload_id) VALUES (%s, %s, %s, %s, %s)",
        (row["first_author"], row["publication_year"], row.get("doi"), row.get("title"), upload_id),
    )
    return cur.lastrowid


def _source_ids(cur, row, upload_id):
    if row["source"] == "sequencing":
        return None, _batch_id(cur, row["state_id"], row["year"])
    return _publication_id(cur, row, upload_id), None


def commit(cur, kind, rows, upload_id):
    if kind == "who_thresholds":
        return commit_rules(cur, rows)
    inserted = 0
    for row in rows:
        if kind in ("sequencing_mutations", "publication_mutations"):
            pub_id, batch_id = _source_ids(cur, row, upload_id)
            cur.execute(
                """
                INSERT INTO observations (mutation_id, state_id, year, source_type, publication_id, sequencing_batch_id,
                                          city, prevalence, sample_count, upload_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    _mutation_id(cur, row["gene"], row["mutation"]),
                    row["state_id"], row["year"], row["source"], pub_id, batch_id,
                    row.get("city"), row["prevalence"], row.get("sample_count"), upload_id,
                ),
            )
        elif kind == "species":
            pub_id, batch_id = _source_ids(cur, row, upload_id)
            cur.execute(
                """
                INSERT INTO species_observations (state_id, year, source_type, publication_id, sequencing_batch_id,
                                                  species_combination, sample_count, upload_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (row["state_id"], row["year"], row["source"], pub_id, batch_id, row["species"], row["sample_count"], upload_id),
            )
        elif kind == "diagnostic":
            pub_id, batch_id = _source_ids(cur, row, upload_id)
            cur.execute(
                """
                INSERT INTO hrp_deletions (state_id, year, source_type, publication_id, sequencing_batch_id,
                                           deletion_type, location, sample_count, upload_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (row["state_id"], row["year"], row["source"], pub_id, batch_id, row["deletion_type"],
                 row.get("location"), row["sample_count"], upload_id),
            )
        elif kind == "moi":
            batch_id = _batch_id(cur, row["state_id"], row["year"])
            cur.execute(
                "INSERT INTO moi_distribution (sequencing_batch_id, moi_value, sample_count, upload_id) VALUES (%s, %s, %s, %s)",
                (batch_id, row["moi"], row["sample_count"], upload_id),
            )
        elif kind == "mis":
            cur.execute(
                "INSERT INTO mis_prevalence (state_id, year, prevalence, upload_id) VALUES (%s, %s, %s, %s)",
                (row["state_id"], row["year"], row["prevalence"], upload_id),
            )
        inserted += 1
    return inserted


def _rule_pattern(text: str) -> str:
    t = re.sub(r"\s+", " ", str(text or "").strip())
    if not t:
        return ""
    if t.lower().startswith("any additional") or t == "*":
        return "*"
    if t.lower() == "deletion":
        return "deletion"
    return "-".join(p.strip().upper() for p in re.split(r"[/,;]", t) if p.strip())


def _rule_gene(text: str) -> str:
    g = re.sub(r"\s+", "", str(text or "").lower())
    if g in ("dhps/dhfr", "dhfr/dhps", "psfr"):
        return "psfr"
    return GENE_ALIASES.get(g, g)


def validate_rules(filename: str, content: bytes):
    header, body = read_rows(filename, content)
    if sum(1 for h in header if _text(h)) < 3 and body:
        header, body = body[0], body[1:]
    norm = [_norm_header(h) for h in header]
    wanted = {"category": "category", "antimalarial": "antimalarial", "gene": "marker", "mutation": "haplotype",
              "who_status": "who status", "min": "min", "max": "max", "classification": "classification",
              "message": "message", "summary": "action", "comment": "comment"}
    positions = {}
    for key, label in wanted.items():
        for i, h in enumerate(norm):
            if h == label or HEADER_ALIASES.get(h) == label:
                positions[key] = i
                break
    missing = [label for key, label in (("gene", "Marker"), ("mutation", "Haplotype"), ("max", "Max"), ("classification", "Classification")) if key not in positions]
    if missing:
        return {"total_rows": len(body), "valid_rows": 0, "rows": [], "warnings": [],
                "errors": [{"row": 1, "column": "Header", "message": f"Missing column(s): {', '.join(missing)}"}]}
    errors, warnings, rows = [], [], []
    for index, raw in enumerate(body, start=2):
        v = {k: (raw[p] if p < len(raw) else None) for k, p in positions.items()}
        if not any(_text(x) for x in v.values()):
            continue
        gene, pattern = _rule_gene(v.get("gene")), _rule_pattern(v.get("mutation"))
        if not gene or not pattern:
            errors.append({"row": index, "column": "Marker", "message": "Marker and haplotype are required"})
            continue
        try:
            upper = float(_text(v.get("max")).rstrip("%")) if _text(v.get("max")) else None
            if upper is not None and upper > 1:
                upper /= 100
        except ValueError:
            upper = None
        if upper is None or not 0 < upper <= 1:
            errors.append({"row": index, "column": "Max", "message": "Max must be between 0 and 1"})
            continue
        cls = norm_key(v.get("classification"))
        if cls not in ("low", "intermediate", "high"):
            errors.append({"row": index, "column": "Classification", "message": "Use low, intermediate or high"})
            continue
        message = _text(v.get("message")) or None
        status = norm_key(v.get("who_status"))
        if status not in ("validated", "candidate", "none"):
            low = (message or "").lower()
            status = "validated" if "validated mutation" in low else "candidate" if "candidate mutation" in low else "none"
            if pattern == "*":
                status = "none"
        antimalarial = _text(v.get("antimalarial")) or None
        rows.append({"_row": index, "gene": gene, "mutation": pattern, "who_status": status,
                     "antimalarial": None if (antimalarial or "").upper() == "NA" else antimalarial,
                     "max": round(upper, 6), "classification": cls, "message": message,
                     "summary": _text(v.get("summary")) or None, "comment": _text(v.get("comment")) or None})
    groups = {}
    for r in rows:
        groups.setdefault((r["gene"], r["mutation"]), []).append(r)
    clean = []
    for (gene, pattern), levels in groups.items():
        levels.sort(key=lambda r: r["max"])
        maxes = [r["max"] for r in levels]
        if len(set(maxes)) != len(maxes):
            errors.append({"row": levels[0]["_row"], "column": "Max",
                           "message": f"{gene} {pattern}: two levels end at the same value ({', '.join(str(m) for m in maxes)})"})
            continue
        if maxes[-1] < 1:
            warnings.append({"row": levels[-1]["_row"], "column": "Max", "message": f"{gene} {pattern}: last level ends at {maxes[-1]}, values above it use that level"})
        statuses = {r["who_status"] for r in levels}
        if len(statuses) > 1:
            warnings.append({"row": levels[0]["_row"], "column": "WHO status", "message": f"{gene} {pattern}: mixed status, using {sorted(statuses)[-1]}"})
        clean.extend(levels)
    return {"total_rows": len(body), "valid_rows": len(clean), "errors": errors, "warnings": warnings, "rows": clean}


def commit_rules(cur, rows):
    groups = {}
    for r in rows:
        groups.setdefault((r["gene"], r["mutation"]), []).append(r)
    for (gene, pattern), levels in groups.items():
        levels.sort(key=lambda r: r["max"])
        first = levels[0]
        status = "validated" if any(r["who_status"] == "validated" for r in levels) else "candidate" if any(r["who_status"] == "candidate" for r in levels) else "none"
        antimalarial = next((r["antimalarial"] for r in levels if r["antimalarial"]), None)
        reference = next((r["comment"] for r in levels if r["comment"]), None)
        cur.execute("SELECT id FROM alert_rules WHERE gene = %s AND mutation_pattern = %s", (gene, pattern))
        found = cur.fetchone()
        previous = {}
        if found:
            rule_id = found["id"]
            cur.execute("SELECT classification, summary FROM alert_levels WHERE alert_rule_id = %s AND summary IS NOT NULL", (rule_id,))
            previous = {r["classification"]: r["summary"] for r in cur.fetchall()}
            cur.execute("UPDATE alert_rules SET who_status = %s, antimalarial = %s, reference_text = %s WHERE id = %s",
                        (status, antimalarial, reference, rule_id))
            cur.execute("DELETE FROM alert_levels WHERE alert_rule_id = %s", (rule_id,))
        else:
            cur.execute("INSERT INTO alert_rules (gene, mutation_pattern, who_status, antimalarial, reference_text) VALUES (%s, %s, %s, %s, %s)",
                        (gene, pattern, status, antimalarial, reference))
            rule_id = cur.lastrowid
        for order, r in enumerate(levels, start=1):
            cur.execute(
                "INSERT INTO alert_levels (alert_rule_id, level_order, max_prevalence, classification, message, summary) VALUES (%s, %s, %s, %s, %s, %s)",
                (rule_id, order, r["max"], r["classification"], r["message"], r["summary"] or previous.get(r["classification"]) or first["summary"]),
            )
    return len(groups)