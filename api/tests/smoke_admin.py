#!/usr/bin/env python3
# End-to-end check of the admin API against the TEST database only.
# Reset it first with sql/setup_test_database.ps1, then run from the api folder:
#   DB_NAME=malaria_dashboard_test python tests/smoke_admin.py
# Author: Khadim Gueye

import io
import os
import sys

if os.environ.get("DB_NAME") != "malaria_dashboard_test":
    sys.exit("Refusing to run: set DB_NAME=malaria_dashboard_test")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app

client = TestClient(app)
OWNER = "khadimg@run.edu.ng"
results = []


def check(label, condition, detail=""):
    results.append((label, bool(condition)))
    print(("PASS " if condition else "FAIL ") + label + (f"  ({detail})" if detail and not condition else ""))


def auth(token):
    return {"Authorization": f"Bearer {token}"}


r = client.post("/api/auth/login", json={"email": OWNER, "password": "wrong"})
check("wrong password rejected", r.status_code == 401)

r = client.post("/api/auth/login", json={"email": OWNER, "password": "admin"})
check("owner login", r.status_code == 200, r.text)
owner_token = r.json()["token"]
check("owner must change password", r.json()["user"]["must_change_password"])

r = client.get("/api/admin/overview", headers=auth(owner_token))
check("admin area blocked before password change", r.status_code == 403)

r = client.post("/api/auth/change-password", headers=auth(owner_token), json={"current_password": "admin", "new_password": "short"})
check("weak password refused", r.status_code == 400)

r = client.post("/api/auth/change-password", headers=auth(owner_token), json={"current_password": "admin", "new_password": "Owner2026pass"})
check("owner password changed", r.status_code == 200, r.text)

r = client.post("/api/admin/users", headers=auth(owner_token), json={"email": "vera@test.ng", "full_name": "Vera", "role": "admin"})
check("super admin creates admin", r.status_code == 200, r.text)
admin_id, admin_temp = r.json()["id"], r.json()["temporary_password"]

r = client.post("/api/auth/login", json={"email": "vera@test.ng", "password": admin_temp})
admin_token = r.json()["token"]
client.post("/api/auth/change-password", headers=auth(admin_token), json={"current_password": admin_temp, "new_password": "Vera2026pass"})

r = client.get("/api/admin/users", headers=auth(admin_token))
check("admin cannot list users", r.status_code == 403)
r = client.get("/api/admin/audit", headers=auth(admin_token))
check("admin cannot read audit log", r.status_code == 403)

owner_id = client.get("/api/auth/me", headers=auth(owner_token)).json()["id"]
r = client.patch(f"/api/admin/users/{owner_id}", headers=auth(owner_token), json={"role": "admin"})
check("owner cannot be modified", r.status_code in (400, 403))

r = client.get("/api/admin/templates/sequencing_mutations", headers=auth(admin_token))
check("template download", r.status_code == 200 and r.content[:2] == b"PK")
wb = load_workbook(io.BytesIO(r.content))
check("template sheets", wb.sheetnames == ["Instructions", "Data", "Example", "Lists"], wb.sheetnames)

data = wb["Data"]
for col, value in enumerate([2026, "Kano", "kelch13", "TESTSMOKE1", "12%", 100], start=1):
    data.cell(row=2, column=col, value=value)
for col, value in enumerate([2026, "Nowhere", "kelch13", "TESTSMOKE2", 0.2, 50], start=1):
    data.cell(row=3, column=col, value=value)
buffer = io.BytesIO()
wb.save(buffer)
bad_file = buffer.getvalue()
files = {"file": ("test.xlsx", bad_file, "application/octet-stream")}
r = client.post("/api/admin/uploads/preview", headers=auth(admin_token), data={"kind": "sequencing_mutations"}, files=files)
check("preview flags unknown state", r.status_code == 200 and r.json()["error_count"] == 1, r.text[:300])

r = client.post("/api/admin/uploads/commit", headers=auth(admin_token), data={"kind": "sequencing_mutations"}, files=files)
check("commit refused while errors remain", r.status_code == 400)

data.delete_rows(3)
buffer = io.BytesIO()
wb.save(buffer)
files = {"file": ("test.xlsx", buffer.getvalue(), "application/octet-stream")}
r = client.post("/api/admin/uploads/commit", headers=auth(admin_token), data={"kind": "sequencing_mutations"}, files=files)
check("valid upload committed", r.status_code == 200 and r.json()["rows_inserted"] == 1, r.text[:300])
upload_id = r.json().get("upload_id")

obs = client.get("/api/observations", params={"mutation": "TESTSMOKE1"}).json()
check("uploaded row visible on public site", len(obs) == 1 and abs(obs[0]["prevalence"] - 0.12) < 1e-9, obs)

csv_content = "Year,State,Prevalence\n2026,FCT,18.5%\n".encode()
r = client.post("/api/admin/uploads/commit", headers=auth(admin_token), data={"kind": "mis"}, files={"file": ("mis.csv", csv_content, "text/csv")})
check("csv upload with state alias", r.status_code == 200, r.text[:300])

r = client.post(f"/api/admin/uploads/{upload_id}/hide", headers=auth(admin_token))
check("admin hides upload", r.status_code == 200)
obs = client.get("/api/observations", params={"mutation": "TESTSMOKE1"}).json()
check("hidden upload removed from public site", len(obs) == 0)

r = client.post(f"/api/admin/uploads/{upload_id}/restore", headers=auth(admin_token))
check("admin cannot restore", r.status_code == 403)
r = client.post(f"/api/admin/uploads/{upload_id}/restore", headers=auth(owner_token))
check("super admin restores", r.status_code == 200)
obs = client.get("/api/observations", params={"mutation": "TESTSMOKE1"}).json()
check("restored upload visible again", len(obs) == 1)

rules = client.get("/api/admin/alerts", headers=auth(admin_token)).json()
rule = next(x for x in rules if len(x["levels"]) == 3)
levels = [
    {"id": l["id"], "max_prevalence": float(l["max_prevalence"]), "classification": l["classification"],
     "message": l["message"], "summary": l["summary"], "guideline": l["guideline"]}
    for l in rule["levels"]
]
bad = [dict(l) for l in levels]
bad[0]["max_prevalence"] = 0.9
r = client.put(f"/api/admin/alerts/{rule['id']}", headers=auth(admin_token), json={"antimalarial": rule["antimalarial"], "reference_text": rule["reference_text"], "levels": bad})
check("non increasing thresholds refused", r.status_code == 400)
levels[0]["summary"] = "Edited by smoke test"
r = client.put(f"/api/admin/alerts/{rule['id']}", headers=auth(admin_token), json={"antimalarial": rule["antimalarial"], "reference_text": rule["reference_text"], "levels": levels})
check("threshold message updated", r.status_code == 200, r.text[:300])

r = client.put("/api/admin/settings", headers=auth(admin_token), json={"theme.font_scale": 1.1, "site.announcement": "Test banner"})
check("settings updated", r.status_code == 200 and r.json()["theme.font_scale"] == 1.1)
r = client.put("/api/admin/settings", headers=auth(admin_token), json={"theme.primary": "blue"})
check("invalid colour refused", r.status_code == 400)
check("public settings reflect change", client.get("/api/settings").json()["site.announcement"] == "Test banner")

r = client.patch(f"/api/admin/users/{admin_id}", headers=auth(owner_token), json={"role": "super_admin"})
check("super admin promotes admin", r.status_code == 200)
r = client.patch(f"/api/admin/users/{admin_id}", headers=auth(owner_token), json={"is_active": False})
check("super admin removes admin", r.status_code == 200)
r = client.get("/api/admin/overview", headers=auth(admin_token))
check("removed admin locked out", r.status_code == 401)

log = client.get("/api/admin/audit", headers=auth(owner_token)).json()
actions = set(log["actions"])
expected = {"login", "login_failed", "password_changed", "user_created", "data_uploaded", "upload_hidden",
            "upload_restored", "threshold_updated", "settings_updated", "user_role_changed", "user_removed"}
check("audit log records every action", expected <= actions, sorted(expected - actions))

for _ in range(5):
    client.post("/api/auth/login", json={"email": OWNER, "password": "nope"})
r = client.post("/api/auth/login", json={"email": OWNER, "password": "Owner2026pass"})
check("account locked after 5 failures", r.status_code == 423)

variants = (
    "Year,State,Gene,Mutation,Prevalence,Sample count\n"
    "2024,KANO,KELCH13,c580y,1%,10\n"
    "2024,kano state,PfK13,C580Y,0.01,10\n"
    "2024,Akwa-Ibom,pfcrt,k76t,50%,10\n"
    "2024,fct abuja,Kelch 13,c580y,0.02,10\n"
    "2024,CrossRiver,mdr1,n86y,0.3,10\n"
    "2024,  lagos  ,crt,K76T,0.4,10\n"
).encode()
r = client.post("/api/admin/users", headers=auth(owner_token), json={"email": "case@test.ng", "role": "admin"})
case_temp = r.json().get("temporary_password", "")
case_token = client.post("/api/auth/login", json={"email": "case@test.ng", "password": case_temp}).json().get("token", "")
client.post("/api/auth/change-password", headers=auth(case_token), json={"current_password": case_temp, "new_password": "Case2026pass"})
r = client.post("/api/admin/uploads/preview", headers=auth(case_token), data={"kind": "sequencing_mutations"}, files={"file": ("v.csv", variants, "text/csv")})
body = r.json() if r.status_code == 200 else {}
check("case and spelling variants accepted", body.get("error_count") == 0 and body.get("valid_rows") == 6, r.text[:400])
sample = body.get("sample", [])
check("states normalised", [s["state_name"] for s in sample] == ["Kano", "Kano", "Akwa Ibom", "Federal Capital Territory", "Cross River", "Lagos"], [s.get("state_name") for s in sample])
check("genes and mutations normalised", [(s["gene"], s["mutation"]) for s in sample][:4] == [("kelch13", "C580Y"), ("kelch13", "C580Y"), ("crt", "K76T"), ("kelch13", "C580Y")], [(s.get("gene"), s.get("mutation")) for s in sample])
hrp_csv = "Year,State,Source,Deletion type,Sample count\n2024,BORNO,SEQUENCING,PfHRP2,3\n2024,borno,sequencing,HRP2+HRP3,1\n2024,Borno,Sequencing,NONE,40\n".encode()
r = client.post("/api/admin/uploads/preview", headers=auth(case_token), data={"kind": "diagnostic"}, files={"file": ("h.csv", hrp_csv, "text/csv")})
body = r.json() if r.status_code == 200 else {}
check("diagnostic variants accepted", body.get("error_count") == 0 and [s["deletion_type"] for s in body.get("sample", [])] == ["hrp2", "dual", "none"], r.text[:400])
species_csv = "Year,State,Source,Species,Sample count\n2024,Kano,sequencing,P. Falciparum + P. MALARIAE,4\n2024,Kano,Sequencing,PF,20\n".encode()
r = client.post("/api/admin/uploads/preview", headers=auth(case_token), data={"kind": "species"}, files={"file": ("s.csv", species_csv, "text/csv")})
body = r.json() if r.status_code == 200 else {}
check("species variants accepted", body.get("error_count") == 0 and [s["species"] for s in body.get("sample", [])] == ["pf,pm", "pf"], r.text[:400])

png = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360f8cfc0f01f0005000201e2211bc30000000049454e44ae426082"
)
r = client.post("/api/admin/media", headers=auth(case_token), files={"file": ("banner.png", png, "image/png")})
media_url = r.json().get("url", "") if r.status_code == 200 else ""
check("banner image uploaded", media_url.startswith("/api/media/image_"), r.text[:200])
check("uploaded image served", client.get(media_url).status_code == 200 if media_url else False)
r = client.post("/api/admin/media", headers=auth(case_token), files={"file": ("virus.exe", b"MZ", "application/octet-stream")})
check("unsafe file type refused", r.status_code == 400)
r = client.put("/api/admin/settings", headers=auth(case_token), json={"hero.mode": "image", "hero.media_url": media_url, "hero.overlay": 0.5, "size.map": 320, "theme.status_high": "#aa00aa"})
check("background and sizes saved", r.status_code == 200 and r.json().get("hero.mode") == "image", r.text[:200])
r = client.put("/api/admin/settings", headers=auth(case_token), json={"hero.media_url": "https://evil.example/x.png"})
check("external media url refused", r.status_code == 400)
r = client.put("/api/admin/settings", headers=auth(case_token), json={"size.map": 5000})
check("map size out of range refused", r.status_code == 400)
if media_url:
    os.remove(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "media", media_url.rsplit("/", 1)[1]))

r = client.post("/api/admin/alerts", headers=auth(case_token), json={"gene": "kelch13", "mutation": "k189t"})
new_rule = r.json().get("id") if r.status_code == 200 else None
check("specific rule created for a haplotype", new_rule is not None, r.text[:200])
public_rule = next((x for x in client.get("/api/alerts").json() if x["gene"] == "kelch13" and x["mutation_pattern"] == "K189T"), None)
check("specific rule copied from general rule", public_rule is not None and len(public_rule["levels"]) == 3, public_rule)
r = client.post("/api/admin/alerts", headers=auth(case_token), json={"gene": "kelch13", "mutation": "K189T"})
check("duplicate specific rule refused", r.status_code == 409)
general = next(x for x in client.get("/api/alerts").json() if x["gene"] == "kelch13" and x["mutation_pattern"] == "*")
r = client.delete(f"/api/admin/alerts/{general['id']}", headers=auth(case_token))
check("general rule cannot be removed", r.status_code == 400)
r = client.delete(f"/api/admin/alerts/{new_rule}", headers=auth(case_token)) if new_rule else None
check("specific rule removed", r is not None and r.status_code == 200)

r = client.put("/api/admin/settings", headers=auth(case_token), json={"site.github_url": "https://github.com/AnieboGenomicsLab", "site.website_url": "https://ighresearch.org/en/"})
check("footer links saved", r.status_code == 200 and r.json().get("site.website_url") == "https://ighresearch.org/en/", r.text[:200])
r = client.put("/api/admin/settings", headers=auth(case_token), json={"site.github_url": "javascript:alert(1)"})
check("unsafe footer link refused", r.status_code == 400)

r = client.put("/api/admin/settings", headers=auth(case_token), json={"site.default_year": "auto"})
check("default year automatic", r.status_code == 200 and r.json().get("site.default_year") == "auto", r.text[:200])
r = client.put("/api/admin/settings", headers=auth(case_token), json={"site.default_year": 2021})
check("default year fixed", r.status_code == 200 and r.json().get("site.default_year") == 2021, r.text[:200])
r = client.put("/api/admin/settings", headers=auth(case_token), json={"site.default_year": 1500})
check("implausible default year refused", r.status_code == 400)

team = client.get("/api/team").json()
check("team seeded with a lab lead", len(team) >= 4 and team[0]["is_lead"] == 1, team[:1])
r = client.post("/api/admin/team", headers=auth(case_token), json={"name": "Test Member", "title": "Tester", "photo_url": "https://evil.example/x.png"})
check("external team photo refused", r.status_code == 400)
r = client.post("/api/admin/team", headers=auth(case_token), json={"name": "Test Member", "title": "Tester", "sort_order": 9})
member_id = r.json().get("id") if r.status_code == 200 else None
check("team member added", member_id is not None and any(m["name"] == "Test Member" for m in client.get("/api/team").json()))
r = client.put(f"/api/admin/team/{member_id}", headers=auth(case_token), json={"name": "Test Member", "linkedin_url": "https://evil.example/in/x"})
check("non LinkedIn link refused", r.status_code == 400)
client.put(f"/api/admin/team/{member_id}", headers=auth(case_token), json={"name": "Test Member", "linkedin_url": "https://www.linkedin.com/in/test-member/", "sort_order": 9})
check("LinkedIn link shown publicly", any(m.get("linkedin_url") == "https://www.linkedin.com/in/test-member/" for m in client.get("/api/team").json()))
check("seeded LinkedIn profiles", sum(1 for m in client.get("/api/team").json() if m.get("linkedin_url")) >= 5)
client.post(f"/api/admin/team/{member_id}/hide", headers=auth(case_token))
check("removed member hidden from public page", not any(m["name"] == "Test Member" for m in client.get("/api/team").json()))
r = client.post(f"/api/admin/team/{member_id}/restore", headers=auth(case_token))
check("admin restores a hidden member", r.status_code == 200 and any(m["name"] == "Test Member" for m in client.get("/api/team").json()))
r = client.post(f"/api/admin/team/{member_id}/alumni", headers=auth(case_token))
team_now = client.get("/api/team").json()
check("member moved to alumni, listed last", r.status_code == 200 and team_now[-1]["name"] == "Test Member" and team_now[-1]["is_alumni"] == 1, [m["name"] for m in team_now])
lead_id = next(m["id"] for m in team_now if m["is_lead"])
check("lab lead cannot become alumni", client.post(f"/api/admin/team/{lead_id}/alumni", headers=auth(case_token)).status_code == 400)
client.post(f"/api/admin/team/{member_id}/active", headers=auth(case_token))
check("alumni back to the team", not next(m for m in client.get("/api/team").json() if m["name"] == "Test Member")["is_alumni"])
client.post(f"/api/admin/team/{member_id}/hide", headers=auth(case_token))

from PIL import Image

big = io.BytesIO()
Image.effect_noise((3600, 2400), 90).convert("RGB").save(big, "JPEG", quality=100)
r = client.post("/api/admin/media", headers=auth(case_token), files={"file": ("big_photo.jpg", big.getvalue(), "image/jpeg")})
saved = r.json().get("url", "") if r.status_code == 200 else ""
size = Image.open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "media", saved.rsplit("/", 1)[-1])).size if saved else (0, 0)
check("large photo accepted and resized", len(big.getvalue()) > 5 * 1024 * 1024 and max(size) == 2400, (len(big.getvalue()), r.status_code, size))
r = client.post("/api/admin/media", headers=auth(case_token), files={"file": ("fake.jpg", b"not really an image", "image/jpeg")})
check("fake image refused", r.status_code == 400)
if saved:
    os.remove(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "media", saved.rsplit("/", 1)[-1]))

others = [m["id"] for m in client.get("/api/team").json() if not m["is_lead"]]
r = client.post("/api/admin/team/reorder", headers=auth(case_token), json={"ids": list(reversed(others))})
check("team order saved", r.status_code == 200 and [m["id"] for m in client.get("/api/team").json() if not m["is_lead"]] == list(reversed(others)), r.text)
check("bad team order refused", client.post("/api/admin/team/reorder", headers=auth(case_token), json={"ids": [others[0], others[0]]}).status_code == 400)
r = client.delete(f"/api/admin/team/{member_id}", headers=auth(case_token))
check("admin cannot delete a member", r.status_code == 403)
r = client.post("/api/admin/team", headers=auth(owner_token), json={"name": "Visible Member"})
visible_id = r.json().get("id")
r = client.delete(f"/api/admin/team/{visible_id}", headers=auth(owner_token))
check("visible member cannot be deleted directly", r.status_code == 400)
client.post(f"/api/admin/team/{visible_id}/hide", headers=auth(owner_token))
r = client.delete(f"/api/admin/team/{visible_id}", headers=auth(owner_token))
kept = next((m for m in client.get("/api/admin/team", headers=auth(owner_token)).json() if m["id"] == visible_id), None)
check("deleted member kept for the record", r.status_code == 200 and kept is not None and kept["is_deleted"] == 1 and kept["deleted_at"], kept)
check("admin cannot restore a deleted member", client.post(f"/api/admin/team/{visible_id}/restore", headers=auth(case_token)).status_code == 403)
r = client.post(f"/api/admin/team/{visible_id}/restore", headers=auth(owner_token))
check("super admin restores a deleted member", r.status_code == 200 and any(m["id"] == visible_id for m in client.get("/api/team").json()))

c580y = next(x for x in client.get("/api/alerts").json() if x["mutation_pattern"] == "C580Y")
p441l = next(x for x in client.get("/api/alerts").json() if x["mutation_pattern"] == "P441L")
check("WHO status seeded", c580y["who_status"] == "validated" and p441l["who_status"] == "candidate", (c580y["who_status"], p441l["who_status"]))

info = open(os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "assets", "information_for_report.xlsx"), "rb").read()
files = {"file": ("information_for_report.xlsx", info, "application/octet-stream")}
r = client.post("/api/admin/uploads/preview", headers=auth(case_token), data={"kind": "who_thresholds"}, files=files)
body = r.json() if r.status_code == 200 else {}
check("original thresholds file read", body.get("valid_rows", 0) > 90 and body.get("error_count") == 1, r.text[:300])
fixed = io.BytesIO()
wb_info = load_workbook(io.BytesIO(info))
ws = wb_info.worksheets[0]
for row in ws.iter_rows(min_row=1):
    if str(row[3].value or "") == "P441L" and row[6].value == "high" and row[4].value == 0 and row[5].value == 1:
        ws.delete_rows(row[0].row)
        break
wb_info.save(fixed)
files = {"file": ("information_for_report.xlsx", fixed.getvalue(), "application/octet-stream")}
r = client.post("/api/admin/uploads/commit", headers=auth(case_token), data={"kind": "who_thresholds"}, files=files)
check("thresholds file imported", r.status_code == 200 and r.json().get("rows_inserted", 0) >= 40, r.text[:300])
rules = client.get("/api/alerts").json()
a724e = next((x for x in rules if x["mutation_pattern"] == "A724E"), None)
p441l = next(x for x in rules if x["mutation_pattern"] == "P441L")
check("new rule created from file", a724e is not None and a724e["levels"][0]["classification"] == "high")
check("candidate rule levels from file", p441l["who_status"] == "candidate" and len(p441l["levels"]) == 2 and p441l["levels"][0]["summary"], p441l)

import security

security.TRUST_PROXY = True
lagos = {"X-Forwarded-For": "102.89.1.1, 10.0.0.1"}
visitor = "smoke-visitor-0001"
r = client.post("/api/track", headers=lagos, json={"visitor_id": visitor, "path": "/dashboard?year=2023"})
client.post("/api/track", headers=lagos, json={"visitor_id": visitor, "path": "/dashboard?year=2023"})
client.post("/api/track", headers=lagos, json={"visitor_id": visitor, "path": "/admin/users"})
client.post("/api/track", headers={**lagos, "User-Agent": "Googlebot/2.1"}, json={"visitor_id": "bot-visitor-01", "path": "/"})
client.post("/api/track", headers={"X-Forwarded-For": "41.82.1.1"}, json={"visitor_id": "smoke-visitor-0002", "path": "/team"})
check("visit recorded", r.status_code == 204)
r = client.get("/api/admin/visitors", headers=auth(case_token))
v = r.json() if r.status_code == 200 else {}
check("visits deduplicated, admin pages and bots ignored", v.get("total", {}).get("visits") == 2, v.get("total"))
check("visitor located by IP", {p["country_code"] for p in v.get("points", [])} == {"NG", "SN"} and any(p["city"] and p["city"].startswith("Lagos") for p in v.get("points", [])), v.get("points"))
check("top pages without query string", any(p["path"] == "/dashboard" for p in v.get("pages", [])), v.get("pages"))
check("visitors map needs login", client.get("/api/admin/visitors").status_code == 401)

r = client.post("/api/report-requests", headers=lagos, json={"name": "A", "email": "bad"})
check("report request validated", r.status_code == 400)
r = client.post("/api/report-requests", headers=lagos, json={"name": "Ada Obi", "email": "ada@health.gov.ng"})
check("download refused without privacy consent", r.status_code == 400 and "privacy" in r.text.lower())
r = client.post("/api/report-requests", headers=lagos, json={"name": "Ada Obi", "email": "Ada@Health.gov.ng", "organization": "NMEP", "page": "/dashboard?year=2023", "consent": True})
check("report request stored", r.status_code == 200)
r = client.get("/api/admin/report-downloads", headers=auth(case_token), params={"search": "nmep"})
d = r.json() if r.status_code == 200 else {}
row = (d.get("rows") or [{}])[0]
check("report download with country from IP", d.get("total") == 1 and row.get("country_code") == "NG" and row.get("email") == "ada@health.gov.ng", d)
check("report download keeps no IP address", "ip_address" not in row and row.get("report") == "Nigeria, 2023", row)
r = client.get("/api/admin/report-downloads.csv", headers=auth(case_token))
check("report downloads CSV export", r.status_code == 200 and "ada@health.gov.ng" in r.text and "ip_address" not in r.text.splitlines()[0] and "/dashboard?year=2023" in r.text)
check("report downloads need login", client.get("/api/admin/report-downloads").status_code == 401)
r = client.post("/api/report-requests", headers=lagos, json={"name": "Ada Obi", "email": "ada@health.gov.ng", "kind": "data", "label": "crt K76T, 2021", "page": "/map", "consent": True})
rows = client.get("/api/admin/report-downloads", headers=auth(case_token)).json().get("rows", [])
check("data download recorded with its type", r.status_code == 200 and rows[0].get("kind") == "data" and rows[0].get("report") == "crt K76T, 2021", rows[:1])
check("unknown download type refused", client.post("/api/report-requests", headers=lagos, json={"name": "Ada Obi", "email": "ada@health.gov.ng", "kind": "x", "consent": True}).status_code == 400)
security.TRUST_PROXY = False

check("seeded activities public", len(client.get("/api/activities").json()) == 4)
r = client.post("/api/admin/activities", headers=auth(case_token), json={"title": "Bad link", "link_url": "javascript:alert(1)"})
check("unsafe activity link refused", r.status_code == 400)
r = client.post("/api/admin/activities", headers=auth(case_token), json={"title": "Training in Kano", "category": "Training", "activity_date": "2026-08-12", "image_url": "/home/lab_3.jpg", "is_featured": True})
act_id = r.json().get("id")
check("activity added", any(a["title"] == "Training in Kano" for a in client.get("/api/activities").json()))
client.post(f"/api/admin/activities/{act_id}/hide", headers=auth(case_token))
check("removed activity hidden", not any(a["id"] == act_id for a in client.get("/api/activities").json()))
check("admin cannot delete activity", client.delete(f"/api/admin/activities/{act_id}", headers=auth(case_token)).status_code == 403)
check("admin restores an archived activity", client.post(f"/api/admin/activities/{act_id}/restore", headers=auth(case_token)).status_code == 200 and any(a["id"] == act_id for a in client.get("/api/activities").json()))
client.post(f"/api/admin/activities/{act_id}/hide", headers=auth(case_token))
check("super admin deletes activity", client.delete(f"/api/admin/activities/{act_id}", headers=auth(owner_token)).status_code == 200)
kept = next((a for a in client.get("/api/admin/activities", headers=auth(owner_token)).json() if a["id"] == act_id), None)
check("deleted activity kept for the record", kept is not None and kept["is_deleted"] == 1 and not any(a["id"] == act_id for a in client.get("/api/activities").json()), kept)
check("super admin restores a deleted activity", client.post(f"/api/admin/activities/{act_id}/restore", headers=auth(owner_token)).status_code == 200 and any(a["id"] == act_id for a in client.get("/api/activities").json()))
order = [a["id"] for a in client.get("/api/activities").json()]
r = client.post("/api/admin/activities/reorder", headers=auth(case_token), json={"ids": list(reversed(order))})
check("activity order saved", r.status_code == 200 and [a["id"] for a in client.get("/api/activities").json()] == list(reversed(order)), r.text)

import shutil
import tempfile

import gene_flow

real_flow = os.path.join(gene_flow.DATA_DIR, "2021")
gene_flow.DATA_DIR = tempfile.mkdtemp(prefix="gene_flow_")
shutil.copytree(real_flow, os.path.join(gene_flow.DATA_DIR, "2021"))
check("gene flow years listed", client.get("/api/genomics/gene-flow/years").json() == [2021])
net = client.get("/api/genomics/gene-flow/2021").json()
top = max(net["nodes"], key=lambda n: n["outdegree"])
check("gene flow network built", net["matched_tips"] == net["tips"] and len(net["nodes"]) == 37 and top["id"] == "Katsina", (net["tips"], top))
check("unknown gene flow year is 404", client.get("/api/genomics/gene-flow/2019").status_code == 404)
tree_bytes = open(os.path.join(real_flow, gene_flow.TREE_FILE), "rb").read()
meta_bytes = open(os.path.join(real_flow, gene_flow.META_FILE), "rb").read()
r = client.post("/api/admin/genomics/gene-flow", headers=auth(case_token), data={"year": "2022"},
                files={"tree": ("t.treefile", b"(A:1,B:1);", "text/plain"), "metadata": ("m.csv", b"Accession,state\nX,Lagos\n", "text/csv")})
check("tree without matching samples refused", r.status_code == 400 and "match" in r.json()["detail"], r.text)
r = client.post("/api/admin/genomics/gene-flow", headers=auth(case_token), data={"year": "2022"},
                files={"tree": ("t.treefile", tree_bytes, "text/plain"), "metadata": ("m.csv", meta_bytes, "text/csv")})
check("gene flow year uploaded", r.status_code == 200 and client.get("/api/genomics/gene-flow/years").json() == [2021, 2022], r.text[:200])
check("gene flow upload needs login", client.post("/api/admin/genomics/gene-flow", data={"year": "2023"}).status_code == 401)
client.post("/api/admin/genomics/gene-flow/2022/hide", headers=auth(case_token))
check("hidden gene flow year not public", client.get("/api/genomics/gene-flow/years").json() == [2021])
check("admin cannot restore gene flow", client.post("/api/admin/genomics/gene-flow/2022/restore", headers=auth(case_token)).status_code == 403)
client.post("/api/admin/genomics/gene-flow/2022/restore", headers=auth(owner_token))
check("super admin restores gene flow", client.get("/api/genomics/gene-flow/years").json() == [2021, 2022])
shutil.rmtree(gene_flow.DATA_DIR, ignore_errors=True)

check("projects tab off by default", "projects" in client.get("/api/settings").json().get("nav.hidden", []))
check("seeded partners not public yet", client.get("/api/partners").json() == [])
seeded = client.get("/api/admin/partners", headers=auth(case_token)).json()
check("partners seeded for the admin", {"PATH", "Africa CDC"} <= {p["name"] for p in seeded}, [p["name"] for p in seeded])
path_id = next(p["id"] for p in seeded if p["name"] == "PATH")
client.post(f"/api/admin/partners/{path_id}/visibility", headers=auth(case_token), json={"visible": True})
check("partner switched on", [p["name"] for p in client.get("/api/partners").json()] == ["PATH"])
r = client.post("/api/admin/partners", headers=auth(case_token), json={"kind": "partner", "name": "Bad", "link_url": "javascript:alert(1)"})
check("unsafe partner link refused", r.status_code == 400)
client.delete(f"/api/admin/partners/{path_id}", headers=auth(case_token))
kept = next(p for p in client.get("/api/admin/partners", headers=auth(case_token)).json() if p["id"] == path_id)
check("deleted partner kept and hidden", kept["is_deleted"] == 1 and client.get("/api/partners").json() == [])
client.post(f"/api/admin/partners/{path_id}/restore", headers=auth(case_token))
check("partner restored", not next(p for p in client.get("/api/admin/partners", headers=auth(case_token)).json() if p["id"] == path_id)["is_deleted"])
settings_now = client.get("/api/settings").json()
r = client.put("/api/admin/settings", headers=auth(case_token), json={**settings_now, "nav.hidden": ["nothing"]})
check("unknown menu tab refused", r.status_code == 400)
r = client.put("/api/admin/settings", headers=auth(case_token), json={**settings_now, "nav.hidden": ["team", "projects"]})
check("menu tabs saved", r.status_code == 200 and client.get("/api/settings").json()["nav.hidden"] == ["team", "projects"], r.text[:200])

for key, bad in (("site.doi", "not a doi"), ("site.contact_email", "nobody"), ("site.data_release_date", "2026-13-40")):
    r = client.put("/api/admin/settings", headers=auth(case_token), json={**settings_now, key: bad})
    check(f"invalid {key} refused", r.status_code == 400, r.text[:120])
r = client.put("/api/admin/settings", headers=auth(case_token), json={**settings_now, "site.doi": "10.5281/zenodo.1234567", "site.data_release": "2026.1"})
check("DOI and data release saved", r.status_code == 200 and client.get("/api/settings").json()["site.doi"] == "10.5281/zenodo.1234567")

from db import execute as db_execute

db_execute("UPDATE admin_users SET failed_attempts = 0, locked_until = NULL WHERE email = %s", (OWNER,))
import mailer
sent = []
mailer.configured = lambda: True
mailer.send = lambda to, subject, text: sent.append((to, text))
r = client.post("/api/auth/login", json={"email": OWNER, "password": "Owner2026pass"})
body = r.json() if r.status_code == 200 else {}
check("super admin asked for an email code", body.get("two_factor") is True and "token" not in body and len(sent) == 1, r.text[:200])
code = "".join(ch for ch in (sent[-1][1] if sent else "") if ch.isdigit())[:6]
challenge = body.get("challenge", "")
r = client.post("/api/auth/verify-code", json={"challenge": challenge, "code": "000000" if code != "000000" else "111111"})
check("wrong sign-in code refused", r.status_code == 401)
r = client.post("/api/auth/verify-code", json={"challenge": challenge, "code": code})
check("right sign-in code gives a session", r.status_code == 200 and r.json().get("token"), r.text[:200])
r = client.post("/api/auth/verify-code", json={"challenge": challenge, "code": code})
check("sign-in code works only once", r.status_code == 401)
r = client.post("/api/auth/login", json={"email": OWNER, "password": "Owner2026pass"})
challenge = r.json().get("challenge", "")
code = "".join(ch for ch in (sent[-1][1] if sent else "") if ch.isdigit())[:6]
for _ in range(5):
    client.post("/api/auth/verify-code", json={"challenge": challenge, "code": "999999" if code != "999999" else "888888"})
r = client.post("/api/auth/verify-code", json={"challenge": challenge, "code": code})
check("sign-in code blocked after 5 wrong tries", r.status_code == 401)
r = client.post("/api/auth/login", json={"email": "case@test.ng", "password": "Case2026pass"})
check("plain admins sign in without a code", r.status_code == 200 and r.json().get("token") and len(sent) == 2, r.text[:200])

failed = [label for label, ok in results if not ok]
print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
sys.exit(1 if failed else 0)
